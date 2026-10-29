import csv
import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
DATA = BASE / "data"
EXCEL = BASE / "excel" / "NovaBank_analitica.xlsx"

EXPECTED_ROWS = {
    "dim_date": 183,
    "dim_customer": 500,
    "dim_account": 800,
    "dim_rule": 12,
    "dim_priority": 3,
    "dim_currency": 3,
    "fact_transactions": 49582,
    "fact_payment_sla": 10820,
    "fact_alerts": 377,
    "fact_cases": 59,
    "fact_daily_ops": 366,
    "fact_ledger_daily": 366,
}


def read_csv(name):
    path = DATA / f"{name}.csv"
    assert path.exists(), f"Falta {path}"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def to_int(value):
    return int(value) if value not in ("", None) else None


def to_float(value):
    return float(value) if value not in ("", None) else None


def test_expected_csv_files_exist():
    for name in EXPECTED_ROWS:
        assert (DATA / f"{name}.csv").exists(), name


def test_expected_row_counts():
    for name, expected in EXPECTED_ROWS.items():
        assert len(read_csv(name)) == expected, name


def test_unique_keys_in_dimensions():
    keys = {
        "dim_date": "date_key",
        "dim_customer": "customer_id",
        "dim_account": "account_id",
        "dim_rule": "rule_id",
        "dim_priority": "priority",
        "dim_currency": "currency",
        "fact_transactions": "transaction_id",
        "fact_alerts": "alert_id",
        "fact_cases": "case_id",
    }
    for table, key in keys.items():
        values = [row[key] for row in read_csv(table)]
        assert len(values) == len(set(values)), f"Clave duplicada en {table}.{key}"


def test_transaction_dates_and_currencies_are_valid():
    dates = {row["date_key"] for row in read_csv("dim_date")}
    currencies = {row["currency"] for row in read_csv("dim_currency")}
    for row in read_csv("fact_transactions"):
        assert row["date_key"] in dates
        assert row["currency"] in currencies


def test_transaction_amounts_are_non_negative():
    for row in read_csv("fact_transactions"):
        assert to_float(row["amount"]) >= 0
        assert to_float(row["fee"]) >= 0


def test_alert_references_are_valid():
    customers = {row["customer_id"] for row in read_csv("dim_customer")}
    rules = {row["rule_id"] for row in read_csv("dim_rule")}
    priorities = {row["priority"] for row in read_csv("dim_priority")}
    allowed_status = {"NUEVA", "EN_ANALISIS", "ESCALADA", "CERRADA_FALSO_POSITIVO", "REPORTADA"}
    for row in read_csv("fact_alerts"):
        assert row["customer_id"] in customers
        assert row["rule_id"] in rules
        assert row["priority"] in priorities
        assert row["status"] in allowed_status
        assert 0 <= to_int(row["risk_score"]) <= 100
        assert to_int(row["is_closed"]) in (0, 1)
        assert to_int(row["is_false_positive"]) in (0, 1)


def test_case_dates_and_flags():
    allowed_status = {"ABIERTO", "EN_ANALISIS", "ESCALADO", "CERRADO", "REPORTADO"}
    for row in read_csv("fact_cases"):
        assert row["status"] in allowed_status
        assert to_int(row["alert_count"]) >= 1
        assert to_int(row["ros_reported"]) in (0, 1)
        assert to_int(row["is_closed"]) in (0, 1)
        if row["is_on_time_close"] not in ("", None):
            assert to_int(row["is_on_time_close"]) in (0, 1)


def test_daily_reconciliation_is_balanced():
    ledger = defaultdict(lambda: [0.0, 0.0])
    for row in read_csv("fact_ledger_daily"):
        key = (row["date_key"], row["currency"])
        ledger[key][0] += to_float(row["debit"])
        ledger[key][1] += to_float(row["credit"])
    for key, (debit, credit) in ledger.items():
        assert abs(debit - credit) < 0.01, key


def test_daily_operations_match_transaction_totals():
    daily_ops = sum(to_int(row["ops"]) for row in read_csv("fact_daily_ops"))
    transactions = len(read_csv("fact_transactions"))
    assert daily_ops == transactions


def test_export_report_is_consistent():
    report = json.loads((DATA / "export_report.json").read_text(encoding="utf-8"))
    assert report["rows"] == EXPECTED_ROWS
    assert report["integrity_check"] == "ok"
    assert report["foreign_key_violations"] == 0
    assert report["alert_lifecycle"]["alerts"] == EXPECTED_ROWS["fact_alerts"]
    assert report["cases"] == EXPECTED_ROWS["fact_cases"]


def test_powerquery_scripts_exist():
    scripts = sorted((BASE / "powerquery").glob("*.pq"))
    names = {path.name for path in scripts}
    required = {
        "00_parametros.pq",
        "dim_fecha.pq",
        "dim_cliente.pq",
        "dim_cuenta.pq",
        "dim_regla.pq",
        "dim_prioridad.pq",
        "dim_moneda.pq",
        "fact_transacciones.pq",
        "fact_pagos_local.pq",
        "fact_alertas.pq",
        "fact_casos.pq",
        "fact_diario_ops.pq",
        "fact_libro_diario.pq",
        "consulta_conciliacion.pq",
    }
    assert required <= names
    for path in scripts:
        assert path.read_text(encoding="utf-8").strip()


def test_dax_has_required_measures():
    text = (BASE / "powerbi" / "medidas.dax").read_text(encoding="utf-8")
    measures = {
        "Total Operations",
        "Total Amount",
        "Total Amount PYG",
        "STP Rate (proxy)",
        "Local p95 Latency (ms)",
        "Payments Within SLA %",
        "Total Alerts",
        "Open Alerts",
        "P1 Alerts",
        "P1 Assigned On Time %",
        "False Positive Rate",
        "Total Cases",
        "Cases On Time %",
        "Reconciliation Difference",
        "Active Customers",
    }
    for measure in measures:
        assert re.search(rf"^{re.escape(measure)}\s*=", text, flags=re.MULTILINE), measure


def test_excel_workbook_structure():
    assert EXCEL.exists()
    with zipfile.ZipFile(EXCEL) as archive:
        assert archive.testzip() is None
    wb = load_workbook(EXCEL, data_only=False)
    expected_sheets = [
        "Portada",
        "KPIs",
        "Parámetros",
        "Conciliación",
        "Hechos_Diarios",
        "Libro_Diario",
        "Alertas",
        "Casos",
        "Pagos_Local",
        "Simulador",
        "Reglas",
    ]
    assert wb.sheetnames == expected_sheets
    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True

    for name in (
        "Param_ComisionLocalPct",
        "Param_ComisionIntlUSD",
        "Param_SLA_Instantaneo_ms",
        "Lista_Monedas",
        "Lista_Tipos",
        "Lista_SiNo",
    ):
        assert name in wb.defined_names

    expected_tables = {
        "KPIs": "tblKPIs",
        "Parámetros": "tblTiposCambio",
        "Conciliación": "tblConciliacion",
        "Hechos_Diarios": "tblHechosDiarios",
        "Libro_Diario": "tblLibroDiario",
        "Alertas": "tblAlertas",
        "Casos": "tblCasos",
        "Pagos_Local": "tblPagosLocal",
        "Simulador": "tblEscenarios",
        "Reglas": "tblReglas",
    }
    for sheet, table in expected_tables.items():
        assert table in wb[sheet].tables, (sheet, table)

    expected_rows = {
        "tblKPIs": 17,
        "tblTiposCambio": 3,
        "tblConciliacion": 366,
        "tblHechosDiarios": 366,
        "tblLibroDiario": 366,
        "tblAlertas": 377,
        "tblCasos": 59,
        "tblPagosLocal": 10820,
        "tblEscenarios": 5,
        "tblReglas": 12,
        "tblSLA": 3,
    }
    for sheet, table in expected_tables.items():
        reference = wb[sheet].tables[table].ref
        start_cell, end_cell = reference.split(":")
        start_row = int("".join(char for char in start_cell if char.isdigit()))
        end_row = int("".join(char for char in end_cell if char.isdigit()))
        assert end_row - start_row == expected_rows[table], (table, reference)
    assert wb["Alertas"].cell(row=381, column=1).value is not None
    assert wb["Pagos_Local"].cell(row=10824, column=1).value is not None


def test_excel_formulas_and_validations():
    wb = load_workbook(EXCEL, data_only=False)
    conciliation = wb.worksheets[3]
    assert conciliation["D5"].value.startswith("=SUMIFS(")
    assert "_xlfn.XLOOKUP" in conciliation["F5"].value
    assert conciliation["G5"].value.startswith("=IFERROR(INDEX(")
    assert "MATCH(" in conciliation["G5"].value

    kpis = wb["KPIs"]
    assert kpis["B13"].value.startswith("=PERCENTILE(")
    assert "CUMPLE" in kpis["E7"].value

    simulator = wb["Simulador"]
    assert "Param_ComisionIntlUSD" in simulator["B10"].value
    assert "_xlfn.XLOOKUP" in simulator["B15"].value

    validations = simulator.data_validations.dataValidation
    assert len(validations) >= 5
    list_validations = [dv for dv in validations if dv.type == "list"]
    assert list_validations
    for dv in list_validations:
        assert dv.formula1.startswith("Lista_") or dv.formula1.startswith('"')

    for sheet in ("KPIs", "Conciliación", "Alertas", "Casos", "Pagos_Local"):
        assert len(wb[sheet].conditional_formatting) > 0, sheet


def test_excel_documentation_and_specs_exist():
    for relative in (
        "docs/diccionario-datos.md",
        "docs/excel.md",
        "docs/tablero.md",
        "powerbi/README.md",
        "tableau/calculos.md",
        "powerquery/README.md",
    ):
        assert (BASE / relative).exists(), relative


def test_tableau_has_core_calculations():
    text = (BASE / "tableau" / "calculos.md").read_text(encoding="utf-8")
    for field in (
        "Total Operations",
        "Local p95 Latency (ms)",
        "Payments Within SLA %",
        "False Positive Rate",
        "Cases On Time %",
        "Reconciliation Difference",
    ):
        assert field in text, field
