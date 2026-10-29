"""Construye la matriz de ejecución UAT en Excel con resultados reales."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_ROOT = Path(__file__).resolve().parents[2]
QA_DIR = REPO_ROOT / "08-qa-uat"
sys.path.insert(0, str(QA_DIR))

from uat.cases import CASES  # noqa: E402
from uat.defects import DEFECTS  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True, color="1F3864")
SUB_FONT = Font(size=11, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILLS = {
    "Aprobado": PatternFill("solid", fgColor="C6EFCE"),
    "Fallido": PatternFill("solid", fgColor="FFC7CE"),
    "Pendiente de ejecución manual": PatternFill("solid", fgColor="FFEB9C"),
    "No ejecutable - funcionalidad fuera del prototipo": PatternFill("solid", fgColor="D9D9D9"),
}


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _widths(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def _load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def build(results_path: Path, sql_path: Path, output_path: Path) -> Path:
    results = _load_json(results_path)
    sql = _load_json(sql_path)
    cases = results.get("cases", [])
    if not cases:
        cases = [
            {
                "id": case["id"],
                "tipo": case["tipo"],
                "nivel": case["nivel"],
                "historia_id": case["historia_id"],
                "epica_id": case["epica_id"],
                "titulo": case["titulo"],
                "prioridad": case["prioridad"],
                "automatizacion": case["automatizacion"],
                "estado_ejecucion": "Pendiente de ejecución automatizada"
                if case["automatizacion"].startswith("Automatizado")
                else "Pendiente de ejecución manual",
                "evidencia": "",
                "observaciones": case["observaciones"],
            }
            for case in CASES
        ]

    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary["A1"] = "NovaBank - Matriz de ejecución UAT"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = "Caso ficticio de portafolio. Resultados obtenidos del entorno local; no son resultados de producción."
    summary["A2"].font = SUB_FONT

    metrics = results.get("summary", {})
    rows = [
        ("Fecha de ejecución", results.get("generated_at", "")),
        ("Comando", results.get("command", "")),
        ("Entorno", results.get("environment", "")),
        ("Casos totales", metrics.get("cases_total", len(cases))),
        ("Automatizados", metrics.get("automated_total", 0)),
        ("Aprobados", metrics.get("automated_passed", 0)),
        ("Fallidos", metrics.get("automated_failed", 0)),
        ("Pendientes de ejecución manual", metrics.get("manual_pending", 0)),
        ("No ejecutables en el prototipo", metrics.get("not_implemented", 0)),
        ("Tasa de éxito automatizada", metrics.get("automated_pass_rate", "0 %")),
        ("Pruebas pytest", results.get("pytest", {}).get("total", 0)),
        ("Pruebas pytest aprobadas", results.get("pytest", {}).get("passed", 0)),
        ("Pruebas pytest fallidas", results.get("pytest", {}).get("failed", 0)),
        ("Consultas SQL OK", sql.get("queries_ok", 0)),
        ("Consultas SQL totales", sql.get("queries_total", 0)),
        ("integrity_check", sql.get("integrity_check", "")),
        ("Violaciones de clave foránea", sql.get("foreign_key_violations", "")),
    ]
    summary["A4"] = "Métrica"
    summary["B4"] = "Valor"
    _style_header(summary, 4)
    for index, (label, value) in enumerate(rows, start=5):
        summary.cell(row=index, column=1, value=label).border = BORDER
        cell = summary.cell(row=index, column=2, value=value)
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(summary, {"A": 36, "B": 80})

    notes_start = len(rows) + 7
    summary.cell(row=notes_start, column=1, value="Notas de alcance").font = Font(bold=True, size=12)
    notes = [
        "Los casos manuales no se declaran ejecutados: quedan pendientes de evidencia del autor.",
        "Los casos 'No implementado' describen funcionalidades especificadas pero ausentes en el prototipo.",
        "La validación ISO 20022 es estructural y de negocio; no se usaron XSD oficiales (S-056).",
        "Los resultados de pytest corresponden a la suite completa ejecutada en este equipo.",
        "Las evidencias de herramientas externas (Jira, ADO, Figma, Power BI, Tableau) están en docs/EVIDENCIAS.md.",
    ]
    for offset, note in enumerate(notes, start=1):
        cell = summary.cell(row=notes_start + offset, column=1, value=f"- {note}")
        cell.alignment = Alignment(wrap_text=True)
        summary.merge_cells(
            start_row=notes_start + offset,
            start_column=1,
            end_row=notes_start + offset,
            end_column=2,
        )

    case_sheet = wb.create_sheet("Casos")
    headers = [
        "ID",
        "Tipo",
        "Nivel",
        "Historia",
        "Épica",
        "Título",
        "Prioridad",
        "Automatización",
        "Resultado",
        "Evidencia",
        "Observaciones",
    ]
    case_sheet.append(headers)
    _style_header(case_sheet)
    for case in cases:
        case_sheet.append(
            [
                case["id"],
                case["tipo"],
                case["nivel"],
                case["historia_id"],
                case["epica_id"],
                case["titulo"],
                case["prioridad"],
                case["automatizacion"],
                case["estado_ejecucion"],
                case["evidencia"],
                case["observaciones"],
            ]
        )
    for row in case_sheet.iter_rows(min_row=2, max_row=case_sheet.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        status = row[8].value
        if status in STATUS_FILLS:
            row[8].fill = STATUS_FILLS[status]
    _widths(
        case_sheet,
        {
            "A": 10,
            "B": 11,
            "C": 14,
            "D": 11,
            "E": 9,
            "F": 52,
            "G": 11,
            "H": 22,
            "I": 32,
            "J": 42,
            "K": 48,
        },
    )
    case_sheet.freeze_panes = "A2"
    case_sheet.auto_filter.ref = f"A1:K{case_sheet.max_row}"
    if case_sheet.max_row > 1:
        _table(case_sheet, "TablaCasos", f"A1:K{case_sheet.max_row}")

    matrix = wb.create_sheet("Matriz ejecucion")
    matrix.append(["ID", "Historia", "Tipo", "Prioridad", "Automatización", "Resultado", "Evidencia"])
    _style_header(matrix)
    for case in cases:
        matrix.append(
            [
                case["id"],
                case["historia_id"],
                case["tipo"],
                case["prioridad"],
                case["automatizacion"],
                case["estado_ejecucion"],
                case["evidencia"],
            ]
        )
    for row in matrix.iter_rows(min_row=2, max_row=matrix.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[5].value in STATUS_FILLS:
            row[5].fill = STATUS_FILLS[row[5].value]
    _widths(matrix, {"A": 10, "B": 11, "C": 11, "D": 11, "E": 24, "F": 32, "G": 48})
    matrix.freeze_panes = "A2"
    matrix.auto_filter.ref = f"A1:G{matrix.max_row}"

    coverage = wb.create_sheet("Cobertura")
    coverage.append(["Historia", "Casos", "Total", "Automatizados", "Manuales", "No implementados"])
    _style_header(coverage)
    coverage_path = QA_DIR / "matriz-cobertura.csv"
    if coverage_path.exists():
        with coverage_path.open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                coverage.append(
                    [
                        row["historia_id"],
                        row["casos"],
                        int(row["total_casos"]),
                        int(row["automatizados"]),
                        int(row["manuales"]),
                        int(row["no_implementados"]),
                    ]
                )
    for row in coverage.iter_rows(min_row=2, max_row=coverage.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(coverage, {"A": 12, "B": 28, "C": 9, "D": 15, "E": 10, "F": 18})
    coverage.freeze_panes = "A2"

    defects = wb.create_sheet("Defectos")
    defect_headers = [
        "ID",
        "Título",
        "Severidad",
        "Prioridad",
        "Estado",
        "Caso",
        "Historia",
        "Ambiente",
        "Descripción",
        "Pasos",
        "Esperado",
        "Obtenido",
        "Evidencia",
        "Impacto",
        "Responsable",
    ]
    defects.append(defect_headers)
    _style_header(defects)
    for defect in DEFECTS:
        defects.append(
            [
                defect["id"],
                defect["titulo"],
                defect["severidad"],
                defect["prioridad"],
                defect["estado"],
                defect["caso"],
                defect["historia"],
                defect["ambiente"],
                defect["descripcion"],
                defect["pasos"],
                defect["esperado"],
                defect["obtenido"],
                defect["evidencia"],
                defect["impacto"],
                defect["responsable"],
            ]
        )
    for row in defects.iter_rows(min_row=2, max_row=defects.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value == "Crítica":
            row[2].fill = PatternFill("solid", fgColor="FFC7CE")
        elif row[2].value == "Alta":
            row[2].fill = PatternFill("solid", fgColor="FFEB9C")
    _widths(
        defects,
        {
            "A": 10,
            "B": 46,
            "C": 10,
            "D": 9,
            "E": 26,
            "F": 14,
            "G": 13,
            "H": 28,
            "I": 60,
            "J": 50,
            "K": 40,
            "L": 40,
            "M": 46,
            "N": 42,
            "O": 28,
        },
    )
    defects.freeze_panes = "A2"
    defects.auto_filter.ref = f"A1:O{defects.max_row}"

    pytest_sheet = wb.create_sheet("Resultados pytest")
    pytest_sheet.append(["Archivo", "Prueba", "Estado", "Duración (s)", "Mensaje"])
    _style_header(pytest_sheet)
    for test in results.get("pytest", {}).get("testcases", []):
        pytest_sheet.append(
            [
                test.get("classname", ""),
                test.get("name", ""),
                test.get("status", ""),
                test.get("time", ""),
                test.get("message", ""),
            ]
        )
    for row in pytest_sheet.iter_rows(min_row=2, max_row=pytest_sheet.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value == "passed":
            row[2].fill = STATUS_FILLS["Aprobado"]
        elif row[2].value in {"failed", "error"}:
            row[2].fill = STATUS_FILLS["Fallido"]
    _widths(pytest_sheet, {"A": 48, "B": 52, "C": 10, "D": 12, "E": 60})
    pytest_sheet.freeze_panes = "A2"
    pytest_sheet.auto_filter.ref = f"A1:E{pytest_sheet.max_row}"

    sql_sheet = wb.create_sheet("Controles SQL")
    sql_sheet.append(["Archivo", "Estado", "Filas", "Duración (ms)", "Error"])
    _style_header(sql_sheet)
    for query in sql.get("queries", []):
        sql_sheet.append(
            [
                query.get("file", ""),
                query.get("status", ""),
                query.get("rows", ""),
                query.get("duration_ms", ""),
                query.get("error") or "",
            ]
        )
    for row in sql_sheet.iter_rows(min_row=2, max_row=sql_sheet.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[1].value == "ok":
            row[1].fill = STATUS_FILLS["Aprobado"]
        elif row[1].value == "error":
            row[1].fill = STATUS_FILLS["Fallido"]
    _widths(sql_sheet, {"A": 40, "B": 10, "C": 10, "D": 14, "E": 50})
    sql_sheet.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    build(
        QA_DIR / "resultados" / "resultados-uat.json",
        QA_DIR / "resultados" / "sql-controls.json",
        QA_DIR / "matriz-ejecucion.xlsx",
    )
    print("Workbook generado")
